"""Checks on a generated report draft.

numbers(docx, reg, xlsx)  every registered Excel number re-read from the .xlsx (cached values) formats to the same
                          text; every numeric token in the document (body, tables, header) is a registered number or a
                          registered literal -> no number typed by hand
forbidden(docx)           investment-judgement wording (build.FORBIDDEN) -> list of hits
human_boxes(docx)         the person-written boxes contain only the label and the writing guide (no generated text)
"""
from __future__ import annotations

from pathlib import Path

from src.report.build import FORBIDDEN, HEADER, HUMAN, TOKEN, Reg, fmtv
from src.report.xlsx import Model


def texts(docx: Path) -> list[str]:
    from docx import Document
    d = Document(str(docx))
    out = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                out.append(cell.text)
    for sec in d.sections:
        out += [p.text for p in sec.header.paragraphs]
    return out


def numbers(docx: Path, reg: Reg, xlsx: Path) -> dict:
    m = Model(xlsx)
    bad_values = []
    n_excel = 0
    for e in reg.entries:
        if e["kind"] != "excel":
            continue
        n_excel += 1
        v = m.val(e["ref"])
        if fmtv(v, e["fmt"]) != e["text"]:
            bad_values.append((e["ref"], e["text"], fmtv(v, e["fmt"])))
    allowed = reg.tokens()
    doc_tokens = [t for s in texts(docx) for t in TOKEN.findall(s)]
    stray = sorted({t for t in doc_tokens if t not in allowed})
    return {"n_excel": n_excel, "n_config": sum(e["kind"] == "config" for e in reg.entries), "bad_values": bad_values,
            "n_tokens": len(doc_tokens), "stray_tokens": stray}


def forbidden(docx: Path) -> list[tuple[str, str]]:
    return [(w, s[:80]) for s in texts(docx) for w in FORBIDDEN if w in s]


def human_boxes(docx: Path) -> dict:
    boxes = [s for s in texts(docx) if s.startswith(HUMAN)]
    bad = []
    for b in boxes:
        lines = [x for x in b.split("\n") if x.strip()]
        if len(lines) != 2 or not lines[1].startswith("작성 안내:"):
            bad.append(b[:80])
    return {"n": len(boxes), "not_empty": bad}


def header_on_all_pages(docx: Path) -> bool:
    from docx import Document
    d = Document(str(docx))
    return all(any(p.text == HEADER for p in sec.header.paragraphs) for sec in d.sections)


def josa_errors(docx: Path) -> list[str]:
    """Wrong particle after a number + unit (e.g. '12,345원와')."""
    from src.report.josa import BAD
    return [s[:80] for s in texts(docx) if BAD.search(s)]
