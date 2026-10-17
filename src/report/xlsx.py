"""Read the cached values of a generated model (openpyxl, data_only) by defined name or address, plus the Inputs rows.

The report takes every number from here, so a report number is always an Excel cell value (or a config value, tagged).
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.utils.cell import range_boundaries


class Model:
    def __init__(self, path: Path):
        self.path = Path(path)
        self.wb = openpyxl.load_workbook(self.path, data_only=True)
        self.names = {}
        for nm, dn in self.wb.defined_names.items():
            dest = list(dn.destinations)
            if dest:
                sh, ref = dest[0]
                self.names[nm] = (sh, ref.replace("$", ""))

    def ref(self, name_or_ref: str) -> tuple[str, str]:
        if name_or_ref in self.names:
            return self.names[name_or_ref]
        sh, cell = name_or_ref.split("!")
        return sh, cell.replace("$", "")

    def val(self, name_or_ref: str):
        sh, cell = self.ref(name_or_ref)
        return self.wb[sh][cell].value

    def addr(self, name_or_ref: str) -> str:
        sh, cell = self.ref(name_or_ref)
        return f"{sh}!{cell}"

    def rng(self, name: str) -> list[list]:
        """2-D list of (value, address) for a range name."""
        sh, ref = self.ref(name)
        c0, r0, c1, r1 = range_boundaries(ref)
        ws = self.wb[sh]
        return [[(ws.cell(r, c).value, f"{sh}!{ws.cell(r, c).coordinate}") for c in range(c0, c1 + 1)] for r in range(r0, r1 + 1)]

    def inputs_rows(self) -> dict[str, dict]:
        """Inputs scalar rows by address of the value cell: label, value, unit, source, as-of, status, basis."""
        ws = self.wb["Inputs"]
        out = {}
        for r in range(5, ws.max_row + 1):
            lab = ws.cell(r, 1).value
            st = ws.cell(r, 6).value
            if lab and st in ("✅", "🔸", "⚠️"):
                out[f"Inputs!B{r}"] = {"label": lab, "value": ws.cell(r, 2).value, "unit": ws.cell(r, 3).value or "",
                                       "src": ws.cell(r, 4).value or "", "asof": ws.cell(r, 5).value or "", "status": st,
                                       "basis": ws.cell(r, 7).value or ""}
        return out
